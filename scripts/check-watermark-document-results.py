"""Read API/browser downloads using independent document libraries and compare parts."""
from pathlib import Path
import json
import sys
import zipfile

from docx import Document
from pptx import Presentation
from pypdf import PdfReader
import pikepdf
from PIL import ImageChops

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'apps/worker/python'))
import watermark_file as wm


def unchanged(source, output, changed):
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(output) as b:
        for name in a.namelist():
            if name not in changed: assert a.read(name)==b.read(name), name


def clean_badge(image):
    assert image.convert('RGB').getpixel((280,210))==(24,52,85), 'Watermark/plate remains'
    assert image.convert('RGB').getpixel((40,40))==(230,145,50), 'Unselected body changed'


def verify(root, prefix='result-'):
    for row in json.loads((root/'report.json').read_text()):
        name=row['name'];source=root/name;out=root/(prefix+name)
        if not out.exists(): raise AssertionError('Missing download: '+str(out))
        if name=='compatibility.docx':
            doc=Document(out)
            assert doc.tables[0].cell(1,1).text=='42' and doc.paragraphs[1].runs[0].bold
            assert 'KEEP normal header' in doc.sections[0].header.paragraphs[0].text
            with zipfile.ZipFile(out) as z:
                assert b'WORD SAMPLE' not in z.read('word/header1.xml')
                assert b'CONFIDENTIAL' not in z.read('word/footer1.xml')
            unchanged(source,out,{'word/header1.xml','word/footer1.xml'})
        elif name=='image-compatibility.docx':
            raw=wm.office_load(out);rootxml=wm.read_xml(raw['word/document.xml'])
            alt=rootxml.xpath('//mc:AlternateContent//a:blip/@r:embed',namespaces=wm.NS)
            ordinary=rootxml.xpath('//a:blip[not(ancestor::mc:AlternateContent)]/@r:embed',namespaces=wm.NS)
            assert len(alt)==2 and len(set(alt))==1 and alt[0]!=ordinary[0]
            links=wm.relation_map(raw,'word/document.xml')
            from io import BytesIO
            clean_badge(wm.image_open(BytesIO(raw[links[alt[0]]])))
            assert 'KEEP after image' in '\n'.join(p.text for p in Document(out).paragraphs)
            unchanged(source,out,{'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'})
        elif name in ('grouped.pptx','hidden.pptx'):
            before=Presentation(source);after=Presentation(out)
            assert len(after.slides)==2
            assert not any(s.shape_type==6 for s in after.slides[0].shapes)
            for slide in after.slides:
                assert next(s for s in slide.shapes if s.has_table).table.cell(1,1).text=='42'
                assert next(s for s in slide.shapes if s.has_chart).chart.series[0].values==(3.,7.)
                assert 'KEEP speaker notes' in slide.notes_slide.notes_text_frame.text
            changed={'ppt/slides/slide1.xml'}
            if name=='grouped.pptx':
                from io import BytesIO
                old=[s for s in before.slides[0].shapes if s.shape_type==13]
                new=[s for s in after.slides[0].shapes if s.shape_type==13]
                clean_badge(wm.image_open(BytesIO(new[0].image.blob)))
                assert old[1].image.blob==new[1].image.blob
                for a,b in zip(old,new):
                    assert (a.rotation,a.crop_left,a.crop_right,a.width,a.height)==(b.rotation,b.crop_left,b.crop_right,b.width,b.height)
                changed.update({'ppt/slides/_rels/slide1.xml.rels','[Content_Types].xml'})
            else: assert after.slides[0]._element.get('show')=='0'
            unchanged(source,out,changed)
        elif name=='text-state.pdf':
            before=PdfReader(source);after=PdfReader(out)
            assert len(after.pages)==2 and len(after.outline)==len(before.outline)
            text=after.pages[0].extract_text()
            assert 'PDF SAMPLE' not in text and 'KEEP body text' in text and 'KEEP shared font' in text
            assert after.get_fields()['keep']['/V']=='Kept'
            assert str(after.pages[0]['/Annots'][0].get_object()['/A']['/URI'])=='https://example.com'
            assert after.pages[1].extract_text()==before.pages[1].extract_text()
        elif name=='scan.pdf':
            assert PdfReader(out).pages[0].extract_text()==PdfReader(source).pages[0].extract_text()
            with pikepdf.open(out) as pdf:
                page=pdf.pages[0];cmd=wm.pdf_instructions(page)
                name=next(c.operands[0] for c in cmd if str(c.operator)=='Do')
                clean_badge(pikepdf.PdfImage(page.Resources.XObject[name]).as_pil_image())
        print('PASS read-back',row['name'],'editable content and untouched parts preserved')


if __name__=='__main__':
    verify(Path(sys.argv[1]),sys.argv[2] if len(sys.argv)>2 else 'result-')
