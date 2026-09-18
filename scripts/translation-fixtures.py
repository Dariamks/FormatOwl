from pathlib import Path
import json,sys,zipfile
from PIL import Image,ImageDraw,ImageFont
from reportlab.pdfgen import canvas
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'apps/worker/python'))
from translate_file import font_path
root=Path('.data/translation-fixtures');root.mkdir(parents=True,exist_ok=True)
im=Image.new('RGB',(800,500),'#f4efe3');d=ImageDraw.Draw(im);font=ImageFont.truetype(font_path('Welcome'),34)
d.rectangle((50,40,750,460),outline='#bbc8aa',width=3);d.text((100,90),'Welcome to FormatOwl',font=font,fill='#22331d');d.text((100,180),'Fresh coffee and green tea',font=font,fill='#22331d');d.text((100,290),'Open daily: 9:00 - 18:00',font=font,fill='#22331d');im.save(root/'menu.png')
transparent=im.convert('RGBA');transparent.putalpha(210);transparent.save(root/'transparent.png')
(root/'sample.txt').write_text('Welcome to FormatOwl.\n\nKeep your files organized.\n\nThank you for visiting.\n',encoding='utf-8')
c=canvas.Canvas(str(root/'sample.pdf'),pagesize=(600,800));c.setFont('Helvetica',18);c.drawString(50,740,'Welcome to FormatOwl');c.setFont('Helvetica',12);c.drawString(50,690,'Keep your files organized.');c.rect(45,630,300,35);c.drawString(50,645,'Product');c.drawString(210,645,'Price');c.drawString(50,600,'E = mc2');c.showPage();c.drawImage(str(root/'menu.png'),20,100,560,350);c.save()
doc=Document();doc.add_heading('Welcome to FormatOwl',0);p=doc.add_paragraph();p.add_run('Keep your files ').bold=True;p.add_run('organized.');doc.sections[0].header.paragraphs[0].text='FormatOwl guide';table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Product';table.cell(0,1).text='Price';table.cell(1,0).text='Coffee';table.cell(1,1).text='5 USD';doc.add_picture(str(root/'menu.png'),width=__import__('docx').shared.Inches(3));doc.save(root/'sample.docx')
with zipfile.ZipFile(root/'sample.epub','w') as z:
 z.writestr('mimetype','application/epub+zip',compress_type=zipfile.ZIP_STORED)
 z.writestr('META-INF/container.xml','<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
 z.writestr('OEBPS/content.opf','<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">filemorph-test</dc:identifier><dc:title>FormatOwl guide</dc:title><dc:language>en</dc:language></metadata><manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="chapter"/></spine></package>')
 z.writestr('OEBPS/chapter.xhtml','<html xmlns="http://www.w3.org/1999/xhtml"><head><title>FormatOwl guide</title></head><body><h1>Welcome to FormatOwl</h1><p>Keep your <strong>files</strong> organized.</p><p>Thank you for visiting.</p></body></html>')
 z.writestr('OEBPS/nav.xhtml','<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><body><nav epub:type="toc"><ol><li><a href="chapter.xhtml">FormatOwl guide</a></li></ol></nav></body></html>')
(root/'sample.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="800" height="300"><rect width="800" height="300" fill="#f4efe3"/><text x="60" y="120" font-size="36">Welcome to FormatOwl</text></svg>')
print(root)
