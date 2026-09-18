"""Generated public-domain fixtures. No user uploads or private documents."""
import sys,zipfile
from pathlib import Path
from io import BytesIO
from PIL import Image,ImageDraw,ImageFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from docx import Document
from lxml import etree
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'apps/worker/python'))
from watermark_file import NS,REL,CT,xml_bytes
FOLDER=Path(sys.argv[1]) if len(sys.argv)>1 else Path('.data/watermark-test')
FOLDER.mkdir(parents=True,exist_ok=True)
def image_fixture(path):
    im=Image.new('RGBA',(600,400),'#f8f9fa');d=ImageDraw.Draw(im);d.rectangle((35,35,190,100),fill='#263bc4');d.text((50,58),'KEEP THIS CONTENT',fill='white');d.text((385,320),'SAMPLE',fill='#6d727a',font=ImageFont.truetype('apps/worker/assets/NotoSansSC-Regular.ttf',28));im.save(path)
image_fixture(FOLDER/'sample.png')
# Compact rounded backplate with tightly bounded OCR text for detection regression.
im=Image.new('RGB',(360,260),'#183455');d=ImageDraw.Draw(im)
d.rectangle((20,20,70,60),fill='#e69132')
d.rounded_rectangle((258,191,354,235),radius=21,fill='#bfc0c3')
for x in range(272,342,12):
    d.rectangle((x,202,x+5,223),fill='#142743');d.rectangle((x,209,x+8,214),fill='#142743')
im.save(FOLDER/'badge.png')
# Texture and gradient ensure live repair goes through the model, not a flat fill.
im=Image.new('RGB',(512,320));d=ImageDraw.Draw(im)
for x in range(512):d.line((x,0,x,320),fill=(80+x//4,120+x//6,160+x//8))
d.text((180,130),'SAMPLE',fill='#4c5967',font=ImageFont.truetype('apps/worker/assets/NotoSansSC-Regular.ttf',30));im.save(FOLDER/'gradient.png')
c=canvas.Canvas(str(FOLDER/'sample.pdf'),pagesize=(600,800))
for i in range(2):
    c.bookmarkPage('p'+str(i));c.addOutlineEntry('Page '+str(i+1),'p'+str(i),0)
    c.setFont('Helvetica',15);c.drawString(45,745,'Preserved body text page '+str(i+1));c.linkURL('https://example.com',(45,720,240,740))
    c.saveState();c.setFillAlpha(.22);c.setFont('Helvetica',54);c.translate(140,330);c.rotate(35);c.drawString(0,0,'SAMPLE');c.restoreState()
    if i==0:c.acroForm.textfield(name='keep-field',x=50,y=660,width=130,height=24,value='Kept')
    c.showPage()
c.save()
c=canvas.Canvas(str(FOLDER/'scan.pdf'),pagesize=(600,400));c.drawImage(str(FOLDER/'sample.png'),0,0,width=600,height=400)
t=c.beginText(50,342);t.setTextRenderMode(3);t.textOut('KEEP THIS CONTENT');c.drawText(t);c.save()
doc=Document();doc.add_paragraph('Preserve this editable document and its header.');doc.add_paragraph('Do not change this body text.');doc.add_picture(str(FOLDER/'sample.png'))
header=doc.sections[0].header;header.paragraphs[0].text='Keep the normal header';p=header.add_paragraph();r=etree.SubElement(p._p,'{'+NS['w']+'}r');pict=etree.SubElement(r,'{'+NS['w']+'}pict')
shape=etree.SubElement(pict,'{'+NS['v']+'}rect',id='Watermark',style='position:absolute;width:300pt;height:70pt;margin-left:70pt;margin-top:500pt',filled='f',stroked='f')
textbox=etree.SubElement(shape,'{'+NS['v']+'}textbox');content=etree.SubElement(textbox,'{'+NS['w']+'}txbxContent');wp=etree.SubElement(content,'{'+NS['w']+'}p');wr=etree.SubElement(wp,'{'+NS['w']+'}r');pr=etree.SubElement(wr,'{'+NS['w']+'}rPr');etree.SubElement(pr,'{'+NS['w']+'}sz').set('{'+NS['w']+'}val','96');etree.SubElement(pr,'{'+NS['w']+'}color').set('{'+NS['w']+'}val','CCCCCC');etree.SubElement(wr,'{'+NS['w']+'}t').text='SAMPLE'
doc.add_page_break();doc.add_paragraph('Preserved second page');doc.save(FOLDER/'sample.docx')
# A small complete OOXML presentation with a repeated master watermark.
P=NS['p'];A=NS['a'];R=NS['r']
def rels(items):return ('<Relationships xmlns="'+REL+'">'+''.join('<Relationship Id="'+i+'" Type="'+R+'/'+typ+'" Target="'+target+'"/>' for i,typ,target in items)+'</Relationships>').encode()
def shape(id,text,x,y,w,h,colour='000000'):
    return f'<p:sp><p:nvSpPr><p:cNvPr id="{id}" name="{text}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US" sz="3200"><a:solidFill><a:srgbClr val="{colour}"/></a:solidFill></a:rPr><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>'
group='<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
ns=f'xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"'
parts={
'_rels/.rels':rels([('rId1','officeDocument','ppt/presentation.xml')]),
'ppt/presentation.xml':f'<p:presentation {ns}><p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst><p:sldIdLst><p:sldId id="256" r:id="rId2"/><p:sldId id="257" r:id="rId3"/></p:sldIdLst><p:sldSz cx="9144000" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>'.encode(),
'ppt/_rels/presentation.xml.rels':rels([('rId1','slideMaster','slideMasters/slideMaster1.xml'),('rId2','slide','slides/slide1.xml'),('rId3','slide','slides/slide2.xml')]),
'ppt/slideMasters/slideMaster1.xml':f'<p:sldMaster {ns}><p:cSld><p:spTree>{group}{shape(2,"SAMPLE",2800000,3000000,4500000,1000000,"CCCCCC")}</p:spTree></p:cSld><p:clrMap accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" bg1="lt1" bg2="lt2" folHlink="folHlink" hlink="hlink" tx1="dk1" tx2="dk2"/><p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst><p:txStyles/></p:sldMaster>'.encode(),
'ppt/slideMasters/_rels/slideMaster1.xml.rels':rels([('rId1','slideLayout','../slideLayouts/slideLayout1.xml')]),
'ppt/slideLayouts/slideLayout1.xml':f'<p:sldLayout {ns} type="blank"><p:cSld name="Blank"><p:spTree>{group}</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>'.encode(),
'ppt/slideLayouts/_rels/slideLayout1.xml.rels':rels([('rId1','slideMaster','../slideMasters/slideMaster1.xml')]),
}
for i in (1,2):
    parts[f'ppt/slides/slide{i}.xml']=f'<p:sld {ns}><p:cSld><p:spTree>{group}{shape(2,"Preserve editable slide "+str(i),500000,600000,7500000,1000000)}</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sld>'.encode()
    parts[f'ppt/slides/_rels/slide{i}.xml.rels']=rels([('rId1','slideLayout','../slideLayouts/slideLayout1.xml')])
content=f'<Types xmlns="{CT}"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
for part in parts:
    if not part.endswith('.xml'):continue
    typ='presentation.main' if part=='ppt/presentation.xml' else 'slideMaster' if '/slideMasters/' in part else 'slideLayout' if '/slideLayouts/' in part else 'slide'
    content+=f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-officedocument.presentationml.{typ}+xml"/>'
parts['[Content_Types].xml']=(content+'</Types>').encode()
with zipfile.ZipFile(FOLDER/'sample.pptx','w',zipfile.ZIP_DEFLATED) as z:
    for name,value in parts.items():z.writestr(name,value)
print('Watermark fixtures:',FOLDER)
