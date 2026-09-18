"""Public-domain document fixtures made with python-docx, python-pptx and ReportLab."""
from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches as WordInches
from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'apps/worker/python'))
import watermark_file as wm

MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'


def alternate_content(choice, fallback):
    node = etree.Element('{'+MC+'}AlternateContent', nsmap={'mc': MC, 'w14': W14})
    etree.SubElement(node, '{'+MC+'}Choice', Requires='w14').append(choice)
    etree.SubElement(node, '{'+MC+'}Fallback').append(fallback)
    return node


def word_fixture(folder):
    path = folder/'compatibility.docx'
    doc = Document()
    doc.add_heading('KEEP editable heading', 1)
    doc.add_paragraph('KEEP body text, formatting and page breaks.').runs[0].bold = True
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'KEEP table'; table.cell(1, 1).text = '42'
    doc.sections[0].header.paragraphs[0].text = 'KEEP normal header'
    para = doc.sections[0].header.add_paragraph()
    pict = etree.fromstring(f'''<w:pict xmlns:w="{wm.NS['w']}" xmlns:v="{wm.NS['v']}">
      <v:rect id="WordWatermark" style="position:absolute;width:240pt;height:50pt;margin-left:45pt;margin-top:350pt" filled="f" stroked="f">
        <v:textbox><w:txbxContent><w:p><w:r><w:rPr><w:sz w:val="64"/></w:rPr><w:t>WORD SAMPLE</w:t></w:r></w:p></w:txbxContent></v:textbox>
      </v:rect></w:pict>''')
    para.add_run()._r.append(alternate_content(deepcopy(pict), deepcopy(pict)))
    doc.sections[0].footer.paragraphs[0].text = 'CONFIDENTIAL'
    doc.add_page_break(); doc.add_paragraph('KEEP second page')
    doc.save(path)
    return path


def word_image_fixture(folder):
    path = folder/'image-compatibility.docx'
    doc = Document(); doc.add_paragraph('KEEP before image')
    picture = doc.add_picture(str(folder/'badge.png'), width=WordInches(3))
    inline = picture._inline; drawing = inline.getparent(); run = drawing.getparent()
    # Both alternatives refer to the same picture. Editing either branch alone
    # would leave a watermark in consumers that choose the other representation.
    run.replace(drawing, alternate_content(deepcopy(drawing), deepcopy(drawing)))
    doc.add_paragraph('KEEP after image')
    doc.add_picture(str(folder/'badge.png'), width=WordInches(3))
    doc.save(path)
    return path


def powerpoint_fixture(folder, hidden=False):
    path = folder/('hidden.pptx' if hidden else 'grouped.pptx')
    prs = Presentation()
    for index in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if hidden and index == 0: slide._element.set('show', '0')
        title = slide.shapes.add_textbox(Inches(.4), Inches(.2), Inches(7), Inches(.5))
        title.text = f'KEEP slide {index+1}'
        table = slide.shapes.add_table(2, 2, Inches(.4), Inches(1), Inches(3), Inches(1)).table
        table.cell(0, 0).text = 'KEEP table'; table.cell(1, 1).text = '42'
        data = CategoryChartData(); data.categories = ['A', 'B']; data.add_series('KEEP series', (3, 7))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(1), Inches(3), Inches(2), data)
        slide.notes_slide.notes_text_frame.text = 'KEEP speaker notes'
        if index == 0:
            group = slide.shapes.add_group_shape()
            plate = group.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6), Inches(5), Inches(3), Inches(.65))
            plate.fill.solid(); plate.fill.fore_color.rgb = RGBColor(192, 192, 192)
            plate.line.fill.background()
            label = group.shapes.add_textbox(Inches(6.15), Inches(5.08), Inches(2.7), Inches(.45))
            label.text = 'GROUP SAMPLE'; label.text_frame.paragraphs[0].font.size = Pt(22)
            # A second, rotated occurrence shares the same underlying PNG.
        for x in (.5, 4.5):
            pic = slide.shapes.add_picture(str(folder/'badge.png'), Inches(x), Inches(3.1), Inches(3), Inches(1.6))
            pic.crop_left = .03; pic.crop_right = .02
            if x > 4: pic.rotation = 12
    prs.save(path)
    return path


def pdf_fixture(folder):
    path = folder/'text-state.pdf'
    c = canvas.Canvas(str(path), pagesize=(600, 800))
    c.setFont('Helvetica', 14); c.drawString(40, 740, 'KEEP body text')
    c.saveState(); c.setFillColorRGB(.8, .8, .8)
    # A common PDF layout has no per-text q/Q pair. Later text reuses font state.
    c.setFont('Helvetica', 40); c.drawString(100, 330, 'PDF SAMPLE')
    c.drawString(40, 250, 'KEEP shared font'); c.restoreState()
    c.linkURL('https://example.com', (40, 720, 150, 740))
    c.acroForm.textfield(name='keep', x=40, y=650, width=140, height=25, value='Kept')
    c.bookmarkPage('one'); c.addOutlineEntry('KEEP bookmark', 'one')
    c.showPage(); c.setFont('Helvetica', 14); c.drawString(40, 740, 'KEEP second page'); c.save()
    return path


def build(folder):
    folder = Path(folder); folder.mkdir(parents=True, exist_ok=True)
    im = Image.new('RGB', (360, 260), '#183455'); d = ImageDraw.Draw(im)
    d.rectangle((20, 20, 70, 60), fill='#e69132')
    d.rounded_rectangle((258, 191, 354, 235), radius=21, fill='#bfc0c3')
    for x in range(272, 342, 12):
        d.rectangle((x, 202, x+5, 223), fill='#142743'); d.rectangle((x, 209, x+8, 214), fill='#142743')
    im.save(folder/'badge.png')
    scan=canvas.Canvas(str(folder/'scan.pdf'),pagesize=(360,260))
    scan.drawImage(str(folder/'badge.png'),0,0,width=360,height=260)
    text=scan.beginText(20,220);text.setTextRenderMode(3);text.textOut('KEEP searchable scan text');scan.drawText(text);scan.save()
    return [word_fixture(folder), word_image_fixture(folder), powerpoint_fixture(folder),
            powerpoint_fixture(folder, hidden=True), pdf_fixture(folder)]


if __name__ == '__main__':
    for file in build(sys.argv[1] if len(sys.argv)>1 else ROOT/'.data/watermark-documents'):
        print(file)
