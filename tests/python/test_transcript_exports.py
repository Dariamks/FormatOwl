import importlib.util
from pathlib import Path
from docx import Document
from pypdf import PdfReader
module_path=Path(__file__).resolve().parents[2]/'apps/worker/python/export_transcript.py'
spec=importlib.util.spec_from_file_location('transcript_export',module_path)
module=importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)

def fixture():
    return {'durationMs':7200000,'speakers':[{'id':'a','name':'张三 & Alice'}],'segments':[{'id':'1','startMs':1200,'endMs':3500,'text':'修改后的中文。Hello <world> & welcome.','speakerId':'a'},{'id':'2','startMs':3600000,'endMs':3601000,'text':'长录音结尾。','speakerId':'a'}]}

def test_all_exports_use_edited_text(tmp_path):
    for fmt in ['txt','docx','pdf','srt','vtt']:
        path=tmp_path/f'transcript.{fmt}'
        module.export_transcript(fixture(),{'format':fmt,'includeSpeakers':True,'includeTimestamps':True},path)
        assert path.stat().st_size>0
        if fmt=='docx': text='\n'.join(p.text for p in Document(path).paragraphs)
        elif fmt=='pdf': text='\n'.join(p.extract_text() for p in PdfReader(path).pages)
        else: text=path.read_text()
        assert '修改后的中文' in text and '长录音结尾' in text
        assert '张三' in text
        if fmt=='srt': assert '01:00:00,000 --> 01:00:01,000' in text
        if fmt=='vtt': assert text.startswith('WEBVTT\n') and '01:00:00.000 --> 01:00:01.000' in text
        if fmt=='pdf': assert '/FontFile2' in str(PdfReader(path).pages[0]['/Resources']['/Font'].get_object()['/F2+0'].get_object()['/FontDescriptor'].get_object())

def test_pdf_wraps_long_cjk_and_markup(tmp_path):
    data=fixture()
    data['segments'][0]['text']='长句<不执行标签>&。'*400
    output=tmp_path/'long.pdf'
    module.export_transcript(data,{'format':'pdf','includeSpeakers':False,'includeTimestamps':False},output)
    pdf=PdfReader(output)
    assert len(pdf.pages)>=2
    assert '不执行标签' in ''.join(p.extract_text() for p in pdf.pages)

def test_arabic_transcript_pdf_is_readable_and_searchable(tmp_path):
 import pypdfium2 as pdfium
 data={'speakers':[],'segments':[{'startMs':0,'endMs':1000,'text':'مرحبا بالعالم','speakerId':None}]}
 path=tmp_path/'arabic.pdf'
 module.export_transcript(data,{'format':'pdf','includeSpeakers':False,'includeTimestamps':False},path)
 assert 'مرحبا' in pdfium.PdfDocument(str(path))[0].get_textpage().get_text_range()
