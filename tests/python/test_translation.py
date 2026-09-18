import sys,json,zipfile
from pathlib import Path
import pytest
from PIL import Image,ImageDraw
sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'apps/worker/python'))
from translate_file import extract,export,read_xml,safe_zip,block,image_render,restore_region
ROOT=Path(__file__).resolve().parents[2]
FIX=ROOT/'.data/translation-fixtures'

def parsed(tmp_path,name):
 return extract({'source':str(FIX/name),'format':name.split('.')[-1],'folder':str(tmp_path),'sourceLanguage':'en','targetLanguage':'zh'})

def test_structured_roundtrips(tmp_path):
 from docx import Document
 for ext in ['docx','epub','txt']:
  data=parsed(tmp_path,f'sample.{ext}')
  assert data['blocks']
  for b in data['blocks']:b['translatedText']='翻译：'+b['sourceText']
  for mode in ['translated','bilingual','original']:
   output=tmp_path/f'{mode}.{ext}'
   export({'data':data,'folder':str(tmp_path),'source':str(FIX/f'sample.{ext}'),'output':str(output),'options':{'format':ext,'mode':mode}})
   assert output.stat().st_size>0
   if ext=='docx':
    doc=Document(output);assert len(doc.tables)==1;assert len(doc.inline_shapes)==1;assert doc.sections[0].header.paragraphs[0].text
    assert any(r.bold for p in doc.paragraphs for r in p.runs)
   if ext=='epub':
    with zipfile.ZipFile(output) as z:
     assert z.read('mimetype')==b'application/epub+zip';assert z.read('OEBPS/content.opf')==zipfile.ZipFile(FIX/'sample.epub').read('OEBPS/content.opf')
     assert b'<strong>' in z.read('OEBPS/chapter.xhtml')

def test_pdf_background_and_searchable_output(tmp_path):
 from pypdf import PdfReader
 data=parsed(tmp_path,'sample.pdf');assert len(data['pages'])==2;assert data['pages'][1]['needsOcr']
 for b in data['blocks']:b['translatedText']='测试';b['style']['fontSize']=12
 output=tmp_path/'translated.pdf';r=export({'data':data,'folder':str(tmp_path),'source':str(FIX/'sample.pdf'),'output':str(output),'options':{'format':'pdf','mode':'bilingual'}})
 assert not r['overflow'];doc=PdfReader(output);assert len(doc.pages)==4;assert '测试' in doc.pages[1].extract_text()
 assert tuple(doc.pages[1].mediabox)==tuple(PdfReader(FIX/'sample.pdf').pages[0].mediabox)

def test_overflow_is_reported(tmp_path):
 data=parsed(tmp_path,'menu.png');b=block('hello',0,box={'x':100,'y':100,'width':50,'height':10,'angle':0});b['translatedText']='a very long translation which cannot fit into this region';data['blocks']=[b]
 _,errors=image_render(data,tmp_path,0,'translated');assert errors==['b0']

def test_alpha_and_outside_pixels_preserved(tmp_path):
 data=parsed(tmp_path,'transparent.png');original=Image.open(tmp_path/'page-0.png').convert('RGBA');patch=Image.new('RGB',(200,100),'blue');patch.save(tmp_path/'patch.png')
 restore_region({'background':str(tmp_path/'page-0.png'),'original':str(tmp_path/'page-0.png'),'patch':str(tmp_path/'patch.png'),'crop':{'x':20,'y':20,'width':200,'height':100},'box':{'x':60,'y':40,'width':80,'height':30},'output':str(tmp_path/'bg.png')})
 result=Image.open(tmp_path/'bg.png');assert result.getpixel((0,0))==original.getpixel((0,0));assert result.getpixel((65,45))[3]==210;assert result.getchannel('A').tobytes()==original.getchannel('A').tobytes()

def test_arabic_searchable_pdf(tmp_path):
 from pypdf import PdfReader
 data=parsed(tmp_path,'sample.pdf');data['targetLanguage']='ar'
 for b in data['blocks']:b['translatedText']='مرحبا بالعالم';b['style']['fontSize']=14;b['box']['width']=min(500,data['pages'][b['page']]['width']-b['box']['x']);b['box']['height']=70
 output=tmp_path/'arabic.pdf';r=export({'data':data,'folder':str(tmp_path),'source':str(FIX/'sample.pdf'),'output':str(output),'options':{'format':'pdf','mode':'translated'}});assert not r['overflow'];import pypdfium2 as pdfium; assert 'مرحبا' in pdfium.PdfDocument(str(output))[0].get_textpage().get_text_range()

def test_reject_external_svg_and_zip_traversal(tmp_path):
 with pytest.raises(ValueError):read_xml(b'<!DOCTYPE svg [<!ENTITY x SYSTEM "file:///etc/passwd">]><svg>&x;</svg>')
 path=tmp_path/'bad.zip'
 with zipfile.ZipFile(path,'w') as z:z.writestr('../bad','hello')
 with pytest.raises(ValueError):safe_zip(path)
 bad=tmp_path/'bad.svg';bad.write_text('<svg xmlns="http://www.w3.org/2000/svg"><image href="http://127.0.0.1/private"/></svg>')
 with pytest.raises(ValueError):extract({'source':str(bad),'format':'svg','folder':str(tmp_path),'sourceLanguage':'en','targetLanguage':'zh'})

def test_rotated_repair_changes_only_polygon(tmp_path):
 from PIL import ImageChops,ImageOps
 from translate_file import region_mask,main
 data=parsed(tmp_path,'transparent.png');source=tmp_path/'page-0.png';box={'x':100,'y':140,'width':200,'height':35,'angle':35}
 crop=main({'action':'crop','source':str(source),'box':box,'output':str(tmp_path/'crop.png')})
 Image.new('RGB',(500,300),'blue').save(tmp_path/'patch.png')
 restore_region({'background':str(source),'original':str(source),'patch':str(tmp_path/'patch.png'),'crop':crop,'box':box,'output':str(tmp_path/'bg.png')})
 a=Image.open(source);b=Image.open(tmp_path/'bg.png');mask=region_mask(a.size,box)
 diff=ImageChops.difference(a,b).convert('RGB');diff.paste((0,0,0),(0,0),mask)
 assert diff.getbbox() is None;assert ImageChops.difference(a,b).convert('RGB').getbbox() is not None

def test_svg_and_bitmap_formats(tmp_path):
 for name in ['sample.svg','menu.png']:
  data=parsed(tmp_path,name)
  for fmt in ['png','jpg','webp']:
   output=tmp_path/f'svg-output.{fmt}';export({'data':data,'folder':str(tmp_path),'source':str(FIX/name),'output':str(output),'options':{'format':fmt,'mode':'original'}})
   with Image.open(output) as im:assert im.width>0

def test_long_text_and_nested_epub_text_order(tmp_path):
 from translate_file import structured_export
 text='hello world '*8000;file=tmp_path/'long.txt';file.write_text(text)
 data=extract({'source':str(file),'format':'txt','folder':str(tmp_path),'sourceLanguage':'en','targetLanguage':'zh'});assert data['blocks'][0]['sourceText']==text
 # Nested emphasis text and tail slots must stay in reading order.
 source=tmp_path/'nested.epub'
 with zipfile.ZipFile(source,'w') as z:z.writestr('chapter.xhtml','<html xmlns="http://www.w3.org/1999/xhtml"><p>A<strong>B<em>C</em>D</strong>E</p></html>')
 data={'format':'epub','blocks':[{'locator':{'part':'chapter.xhtml','index':0},'sourceText':'ABCDE','translatedText':'12345'}]}
 output=tmp_path/'nested-out.epub';structured_export(data,source,output,'translated')
 with zipfile.ZipFile(output) as z:assert ''.join(read_xml(z.read('chapter.xhtml')).itertext())=='12345'


def test_flat_repair_avoids_cloud_and_preserves_crop_size(tmp_path):
 from translate_file import main
 source=tmp_path/'flat.png';crop=tmp_path/'crop.png';output=tmp_path/'patch.png'
 im=Image.new('RGB',(400,240),'white');ImageDraw.Draw(im).text((100,100),'Text',fill='black');im.save(source)
 Image.new('RGB',(384,384),'gray').save(crop)
 result=main({'action':'flat-repair','source':str(source),'crop':str(crop),'box':{'x':90,'y':90,'width':70,'height':30},'output':str(output)})
 assert result['local'] is True
 with Image.open(output) as patch:
  assert patch.size==(384,384);assert patch.getpixel((192,192))==(255,255,255)
 noise=Image.effect_noise((400,240),100).convert('RGB');noise.save(source);output.unlink()
 result=main({'action':'flat-repair','source':str(source),'crop':str(crop),'box':{'x':90,'y':90,'width':70,'height':30},'output':str(output)})
 assert result['local'] is False;assert not output.exists()
