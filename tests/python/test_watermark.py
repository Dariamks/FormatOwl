import sys, json, zipfile, shutil
from pathlib import Path
import pytest
from PIL import Image, ImageChops
from pypdf import PdfReader
sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'apps/worker/python'))
import watermark_file as wm
ROOT=Path(__file__).resolve().parents[2]
FIX=ROOT/'.data/watermark-test'

def analyze(fmt,tmp):
    return wm.analyze({'source':str(FIX/('sample.'+fmt)),'format':fmt,'folder':str(tmp)})
def selection(cid):return {'candidates':[cid],'regions':[]}
def export(fmt,data,sel,tmp,patches=None):
    out=tmp/('output.'+fmt)
    wm.export_file({'source':str(FIX/('sample.'+fmt)),'folder':str(tmp),'data':data,'selection':sel,'patches':patches or {},'output':str(out),'format':fmt,'prefix':'result'})
    return out

def test_image_mask_outside_pixels_and_eraser(tmp_path):
    data=analyze('png',tmp_path);source=FIX/'sample.png'
    r={'id':'1','targetId':'image','kind':'rect','x':.62,'y':.77,'width':.33,'height':.16,'strokes':[{'erase':True,'radius':.01,'points':[[.63,.78]]}]}
    result=wm.prepare_repair({'source':source,'regions':[r],'protected':[],'mask':tmp_path/'mask.png','crop':tmp_path/'crop.png','cropMask':tmp_path/'crop-mask.png','output':tmp_path/'fixed.png'})
    assert result['local']
    orig=wm.image_open(source);fixed=wm.image_open(tmp_path/'fixed.png');mask=Image.open(tmp_path/'mask.png')
    assert fixed.getpixel((378,312))==orig.getpixel((378,312))
    diff=ImageChops.difference(orig,fixed); outside=Image.new('RGBA',orig.size);outside.paste(diff,(0,0),ImageChops.invert(mask));assert not outside.convert('RGB').getbbox()
    assert fixed.getchannel('A').tobytes()==orig.getchannel('A').tobytes()

def test_pdf_preserves_body_forms_links_and_other_page(tmp_path):
    data=analyze('pdf',tmp_path)
    candidates=[c for c in data['candidates'] if c['pages']==[0]]
    assert candidates, data
    c=next(c for c in candidates if 'SAMPLE' in c['label'])
    output=export('pdf',data,selection(c['id']),tmp_path)
    before=PdfReader(FIX/'sample.pdf');after=PdfReader(output)
    assert len(after.pages)==2
    assert 'Preserved body text page 1' in after.pages[0].extract_text()
    assert 'SAMPLE' not in after.pages[0].extract_text()
    assert 'SAMPLE' in after.pages[1].extract_text()
    assert len(after.outline)==len(before.outline)
    assert after.get_fields()['keep-field']['/V']=='Kept'
    assert str(after.pages[0]['/Annots'][0].get_object()['/A']['/URI'])=='https://example.com'

def test_scan_repair_preserves_ocr_layer(tmp_path):
    data=wm.analyze({'source':str(FIX/'scan.pdf'),'format':'pdf','folder':str(tmp_path)})
    assert len(data['targets'])==1
    target=data['targets'][0];patch=wm.image_open(tmp_path/target['file']);patch.paste('#f8f9fa',(375,308,580,380));patch.save(tmp_path/'patched.png')
    out=tmp_path/'scan-output.pdf'
    wm.export_file({'source':str(FIX/'scan.pdf'),'folder':str(tmp_path),'data':data,'selection':selection('missing'),'patches':{target['id']:str(tmp_path/'patched.png')},'output':str(out),'prefix':'scan-result'})
    assert PdfReader(out).pages[0].extract_text()==PdfReader(FIX/'scan.pdf').pages[0].extract_text()

def test_strict_scan_blocks_body_overlap(tmp_path):
    with pytest.raises(ValueError,match='WATERMARK_TEXT_OVERLAP'):
        wm.prepare_repair({'source':FIX/'sample.png','regions':[{'kind':'rect','x':0,'y':0,'width':.5,'height':.5,'strokes':[]}],'protected':[{'x':50,'y':40,'width':100,'height':35}]})

def test_docx_preserves_native_body_and_header(tmp_path):
    data=analyze('docx',tmp_path)
    c=next(c for c in data['candidates'] if 'SAMPLE' in c['label'])
    assert len(c['pages'])==2
    out=export('docx',data,selection(c['id']),tmp_path)
    with zipfile.ZipFile(FIX/'sample.docx') as a,zipfile.ZipFile(out) as b:
        assert a.read('word/document.xml')==b.read('word/document.xml')
        header=b.read(c['locator']['part']);assert b'SAMPLE' not in header and b'Keep the normal header' in header
        for name in a.namelist():
            if name!=c['locator']['part']:assert a.read(name)==b.read(name)

def test_pptx_master_scope_and_native_slides(tmp_path):
    data=analyze('pptx',tmp_path)
    c=next(c for c in data['candidates'] if c['reason']=='master' and 'SAMPLE' in c['label'])
    assert c['pages']==[0,1]
    out=export('pptx',data,selection(c['id']),tmp_path)
    with zipfile.ZipFile(FIX/'sample.pptx') as a,zipfile.ZipFile(out) as b:
        assert a.read('ppt/slides/slide1.xml')==b.read('ppt/slides/slide1.xml')
        assert a.read('ppt/slides/slide2.xml')==b.read('ppt/slides/slide2.xml')
        assert b'SAMPLE' not in b.read(c['locator']['part'])

def test_office_image_repair_keeps_original_media_and_shape(tmp_path):
    data=analyze('docx',tmp_path);target=data['targets'][0];patch=wm.image_open(tmp_path/target['file']);patch.paste('white',(370,300,590,390));patch.save(tmp_path/'fixed.png')
    out=export('docx',data,{'candidates':[],'regions':[]},tmp_path,{target['id']:str(tmp_path/'fixed.png')})
    with zipfile.ZipFile(FIX/'sample.docx') as a,zipfile.ZipFile(out) as b:
        for name in a.namelist():
            if '/media/' in name:assert a.read(name)==b.read(name)
        assert any('/media/fm-' in n for n in b.namelist())
        assert b'Preserve this editable document' in b.read('word/document.xml')

def test_pdf_nested_shared_form_removal_is_per_invocation(tmp_path):
    from reportlab.pdfgen import canvas
    import pikepdf
    source=tmp_path/'shared.pdf';c=canvas.Canvas(str(source),pagesize=(600,800));c.beginForm('shared',0,0,300,80);c.saveState();c.setFont('Helvetica',32);c.drawString(10,30,'SAMPLE');c.restoreState();c.endForm()
    c.saveState();c.translate(20,100);c.doForm('shared');c.restoreState();c.saveState();c.translate(20,300);c.doForm('shared');c.restoreState();c.drawString(20,750,'KEEP');c.save()
    data=wm.analyze({'source':source,'format':'pdf','folder':tmp_path});assert len(data['candidates'])==2
    chosen=data['candidates'][0];output=tmp_path/'out.pdf';wm.export_file({'source':source,'folder':tmp_path,'data':data,'selection':selection(chosen['id']),'output':output,'prefix':'fixed'})
    text=PdfReader(output).pages[0].extract_text();assert text.count('SAMPLE')==1;assert 'KEEP' in text

def test_pdf_shared_image_repair_does_not_change_other_occurrence(tmp_path):
    from reportlab.pdfgen import canvas
    source=tmp_path/'shared-images.pdf';c=canvas.Canvas(str(source),pagesize=(600,800));c.drawImage(str(FIX/'sample.png'),0,0,width=300,height=200);c.drawImage(str(FIX/'sample.png'),300,400,width=300,height=200);c.save()
    data=wm.analyze({'source':source,'format':'pdf','folder':tmp_path});assert len(data['targets'])==2
    target=data['targets'][0];im=wm.image_open(tmp_path/target['file']);im.paste('red',(380,305,585,380));im.save(tmp_path/'patch.png')
    out=tmp_path/'shared-fixed.pdf';wm.export_file({'source':source,'folder':tmp_path,'data':data,'selection':{'candidates':[],'regions':[]},'patches':{target['id']:str(tmp_path/'patch.png')},'output':out,'prefix':'shared-fixed'})
    before=wm.image_open(tmp_path/'page-0.png');after=wm.image_open(tmp_path/'shared-fixed-0.png');diff=ImageChops.difference(before,after).convert('RGB').getbbox();assert diff and diff[2]<450

def test_animation_and_encrypted_pdf_rejected(tmp_path):
    a=Image.new('RGB',(20,20),'red');b=Image.new('RGB',(20,20),'blue');f=tmp_path/'animation.webp';a.save(f,save_all=True,append_images=[b],duration=100)
    with pytest.raises(ValueError,match='UNSUPPORTED_ANIMATION'):wm.image_open(f)
    import pikepdf
    encrypted=tmp_path/'encrypted.pdf'
    with pikepdf.open(FIX/'sample.pdf') as pdf:pdf.save(encrypted,encryption=pikepdf.Encryption(owner='secret',user='secret'))
    with pytest.raises(ValueError,match='ENCRYPTED_PDF'):wm.pdf_check(encrypted)

def test_ai_composite_keeps_alpha_and_pixels_outside_mask(tmp_path):
    source=Image.new('RGBA',(100,100),(40,80,120,100));source.save(tmp_path/'source.png');patch=Image.new('RGBA',(20,20),'red');patch.save(tmp_path/'patch.png');mask=Image.new('L',(100,100));mask.paste(255,(40,40,60,60));mask.save(tmp_path/'mask.png')
    wm.compose({'source':tmp_path/'source.png','mask':tmp_path/'mask.png','patch':tmp_path/'patch.png','box':[30,30,70,70],'output':tmp_path/'out.png'})
    result=wm.image_open(tmp_path/'out.png');assert result.getpixel((0,0))==source.getpixel((0,0));assert result.getpixel((45,45))==(255,0,0,100);assert result.getchannel('A').tobytes()==source.getchannel('A').tobytes()

def test_boundary_correction_reduces_gradient_seam_without_changing_outside(tmp_path):
    from PIL import ImageDraw
    clean=Image.new('RGB',(200,120));generated=Image.new('RGB',clean.size)
    for y in range(clean.height):
        for x in range(clean.width):
            c=(60+x//2,100+x//3,130+x//4);clean.putpixel((x,y),c);generated.putpixel((x,y),(c[0]+10+x//20,c[1]-8,c[2]+5))
    original=clean.copy();ImageDraw.Draw(original).text((80,50),'MARK',fill='black');mask=Image.new('L',clean.size);mask.paste(255,(70,40,140,85))
    result=wm.match_boundary(original,generated.convert('RGBA'),mask)
    errors=[abs(result.getpixel((x,y))[channel]-clean.getpixel((x,y))[channel]) for y in range(40,85) for x in range(70,140) for channel in range(3)]
    assert sum(errors)/len(errors)<1.0

def test_pdf_text_removal_keeps_interleaved_path_state(tmp_path):
    import pikepdf
    with pikepdf.open(FIX/'sample.pdf') as pdf:
        pdf.pages[0].Contents=pdf.make_stream(b'q 10 10 m 50 50 l BT /F1 20 Tf (SAMPLE) Tj ET Q S')
        candidates=wm.pdf_candidates(pdf.pages[0],0)
        assert len(candidates)==1 and candidates[0]['locator']['kind']=='text'
        wm.pdf_edit_container(pdf,pdf.pages[0],[candidates[0]['locator']],[])
        assert b'SAMPLE' not in pdf.pages[0].Contents.read_bytes()
        assert [str(c.operator) for c in wm.pdf_instructions(pdf.pages[0])]==['q','m','l','BT','Tf','Tj','ET','Q','S']

def test_pdf_resource_names_never_replace_an_existing_resource(tmp_path):
    import pikepdf
    source=tmp_path/'collision.pdf'
    with pikepdf.open(FIX/'scan.pdf') as pdf:
        page=pdf.pages[0];cmd=wm.pdf_instructions(page);at=next(i for i,c in enumerate(cmd) if str(c.operator)=='Do');old=page.Resources.XObject[cmd[at].operands[0]]
        name=pikepdf.Name('/FMImage'+str(at));page.Resources.XObject[name]=old
        cmd.append(([name],pikepdf.Operator('Do')));wm.pdf_stream(pdf,page,cmd);pdf.save(source)
    data=wm.analyze({'source':source,'format':'pdf','folder':tmp_path});target=data['targets'][0]
    patch=wm.image_open(tmp_path/target['file']);patch.paste('red',(0,0,20,20));patch.save(tmp_path/'patch.png');out=tmp_path/'out.pdf'
    wm.export_file({'source':source,'folder':tmp_path,'data':data,'selection':{'candidates':[],'regions':[]},'patches':{target['id']:str(tmp_path/'patch.png')},'output':out,'prefix':'fixed'})
    with pikepdf.open(out) as pdf:
        names=[c.operands[0] for c in wm.pdf_instructions(pdf.pages[0]) if str(c.operator)=='Do']
        assert names[0]!=names[-1]
        assert pdf.pages[0].Resources.XObject[names[0]].read_bytes()!=pdf.pages[0].Resources.XObject[names[-1]].read_bytes()

def test_rotated_pdf_keeps_page_rotation_and_text(tmp_path):
    import pikepdf
    source=tmp_path/'rotated.pdf'
    with pikepdf.open(FIX/'sample.pdf') as pdf:pdf.pages[0].Rotate=90;pdf.save(source)
    data=wm.analyze({'source':source,'format':'pdf','folder':tmp_path});c=next(c for c in data['candidates'] if 'SAMPLE' in c['label'] and c['pages']==[0]);assert c['boxes'][0]['width']>0
    out=tmp_path/'out.pdf';wm.export_file({'source':source,'folder':tmp_path,'data':data,'selection':selection(c['id']),'output':out,'prefix':'fixed'})
    result=PdfReader(out);assert result.pages[0].get('/Rotate')==90;assert 'SAMPLE' not in result.pages[0].extract_text();assert 'Preserved body text' in result.pages[0].extract_text()

def test_word_multisection_first_even_and_linked_headers(tmp_path):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from copy import deepcopy
    from lxml import etree
    doc=Document(FIX/'sample.docx');first=doc.sections[0];first.different_first_page_header_footer=True;doc.settings.odd_and_even_pages_header_footer=True
    mark=first.header._element.xpath('.//w:pict')[0].getparent().getparent()
    for header,text in [(first.first_page_header,'FIRST MARK'),(first.even_page_header,'EVEN MARK')]:
        cloned=deepcopy(mark)
        for node in cloned.iter('{'+wm.NS['w']+'}t'):
            if node.text=='SAMPLE':node.text=text
        header._element.append(cloned)
    doc.add_section(WD_SECTION.NEW_PAGE);doc.add_paragraph('SECTION TWO BODY');doc.add_page_break();doc.add_paragraph('SECTION TWO EVEN');doc.add_page_break();doc.add_paragraph('SECTION TWO DEFAULT')
    source=tmp_path/'sections.docx';doc.save(source)
    data=wm.analyze({'source':source,'format':'docx','folder':tmp_path});firstmark=next(c for c in data['candidates'] if 'FIRST MARK' in c['label']);evenmark=next(c for c in data['candidates'] if 'EVEN MARK' in c['label'])
    assert 'section 1 first header' in firstmark['scope'] and 'section 2 first header' in firstmark['scope'];assert 'even header' in evenmark['scope'];assert len(firstmark['pages'])==2
    output=tmp_path/'out.docx';wm.office_apply(source,data,selection(firstmark['id']),{},output)
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(output) as b:
        for name in a.namelist():
            if name!=firstmark['locator']['part']:assert a.read(name)==b.read(name)
        assert b'FIRST MARK' not in b.read(firstmark['locator']['part']);assert b'EVEN MARK' in b.read(evenmark['locator']['part'])

def test_powerpoint_two_masters_remain_independent(tmp_path):
    from lxml import etree
    raw=wm.office_load(FIX/'sample.pptx')
    for old,new in [('ppt/slideMasters/slideMaster1.xml','ppt/slideMasters/slideMaster2.xml'),('ppt/slideMasters/_rels/slideMaster1.xml.rels','ppt/slideMasters/_rels/slideMaster2.xml.rels'),('ppt/slideLayouts/slideLayout1.xml','ppt/slideLayouts/slideLayout2.xml'),('ppt/slideLayouts/_rels/slideLayout1.xml.rels','ppt/slideLayouts/_rels/slideLayout2.xml.rels')]:
        raw[new]=raw[old].replace(b'slideLayout1',b'slideLayout2').replace(b'slideMaster1',b'slideMaster2').replace(b'SAMPLE',b'SECOND MARK')
    raw['ppt/slides/_rels/slide2.xml.rels']=raw['ppt/slides/_rels/slide2.xml.rels'].replace(b'slideLayout1',b'slideLayout2')
    root=wm.read_xml(raw['ppt/presentation.xml']);node=etree.SubElement(root.find('p:sldMasterIdLst',wm.NS),'{'+wm.NS['p']+'}sldMasterId',id='2147483650');node.set(wm.RID+'id','rIdMaster2');raw['ppt/presentation.xml']=wm.xml_bytes(root)
    rels=wm.read_xml(raw['ppt/_rels/presentation.xml.rels']);etree.SubElement(rels,'{'+wm.REL+'}Relationship',Id='rIdMaster2',Type=wm.NS['r']+'/slideMaster',Target='slideMasters/slideMaster2.xml');raw['ppt/_rels/presentation.xml.rels']=wm.xml_bytes(rels)
    ct=wm.read_xml(raw['[Content_Types].xml'])
    for name,typ in [('ppt/slideMasters/slideMaster2.xml','slideMaster'),('ppt/slideLayouts/slideLayout2.xml','slideLayout')]:etree.SubElement(ct,'{'+wm.CT+'}Override',PartName='/'+name,ContentType='application/vnd.openxmlformats-officedocument.presentationml.'+typ+'+xml')
    raw['[Content_Types].xml']=wm.xml_bytes(ct);source=tmp_path/'masters.pptx';wm.write_office(FIX/'sample.pptx',raw,source)
    data=wm.analyze({'source':source,'format':'pptx','folder':tmp_path});first=next(c for c in data['candidates'] if c['label']=='SAMPLE');second=next(c for c in data['candidates'] if c['label']=='SECOND MARK');assert first['pages']==[0];assert second['pages']==[1];assert 'slideLayout1' in first['scope']
    out=tmp_path/'out.pptx';wm.office_apply(source,data,selection(first['id']),{},out)
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(out) as b:
        for name in a.namelist():
            if name!=first['locator']['part']:assert a.read(name)==b.read(name)

def test_logo_on_transparent_background_is_cleared_locally(tmp_path):
    from PIL import ImageDraw
    source=Image.new('RGBA',(160,100),(0,0,0,0));draw=ImageDraw.Draw(source);draw.rectangle((5,5,30,30),fill=(20,80,200,255));draw.ellipse((80,40,105,65),fill=(120,140,160,180));source.save(tmp_path/'source.png')
    result=wm.prepare_repair({'source':tmp_path/'source.png','regions':[{'kind':'rect','x':.45,'y':.3,'width':.3,'height':.45,'strokes':[]}],'protected':[],'mask':tmp_path/'mask.png','crop':tmp_path/'crop.png','cropMask':tmp_path/'small.png','output':tmp_path/'fixed.png'})
    assert result['local'];fixed=wm.image_open(tmp_path/'fixed.png');assert fixed.getpixel((95,50))==(0,0,0,0);assert fixed.crop((0,0,40,40)).tobytes()==source.crop((0,0,40,40)).tobytes()

def test_powerpoint_background_image_repair_keeps_shapes(tmp_path):
    from lxml import etree
    raw=wm.office_load(FIX/'sample.pptx');part='ppt/slides/slide1.xml';root=wm.read_xml(raw[part]);csld=root.find('p:cSld',wm.NS);bg=etree.Element('{'+wm.NS['p']+'}bg');csld.insert(0,bg);prop=etree.SubElement(bg,'{'+wm.NS['p']+'}bgPr');fill=etree.SubElement(prop,'{'+wm.NS['a']+'}blipFill');blip=etree.SubElement(fill,'{'+wm.NS['a']+'}blip');blip.set(wm.RID+'embed','rIdBackground');stretch=etree.SubElement(fill,'{'+wm.NS['a']+'}stretch');etree.SubElement(stretch,'{'+wm.NS['a']+'}fillRect');etree.SubElement(prop,'{'+wm.NS['a']+'}effectLst');raw[part]=wm.xml_bytes(root)
    rp=wm.relpath(part);rels=wm.read_xml(raw[rp]);etree.SubElement(rels,'{'+wm.REL+'}Relationship',Id='rIdBackground',Type=wm.NS['r']+'/image',Target='../media/background.png');raw[rp]=wm.xml_bytes(rels);raw['ppt/media/background.png']=(FIX/'sample.png').read_bytes()
    ct=wm.read_xml(raw['[Content_Types].xml']);etree.SubElement(ct,'{'+wm.CT+'}Default',Extension='png',ContentType='image/png');raw['[Content_Types].xml']=wm.xml_bytes(ct);source=tmp_path/'background.pptx';wm.write_office(FIX/'sample.pptx',raw,source)
    data=wm.analyze({'source':source,'format':'pptx','folder':tmp_path});target=data['targets'][0];assert target['pages']==[0]
    patch=wm.image_open(tmp_path/target['file']);patch.paste('white',(370,300,590,390));patch.save(tmp_path/'patch.png');out=tmp_path/'out.pptx';wm.office_apply(source,data,{'candidates':[],'regions':[]},{target['id']:str(tmp_path/'patch.png')},out)
    with zipfile.ZipFile(out) as z:
        assert z.read('ppt/media/background.png')==raw['ppt/media/background.png'];changed=wm.read_xml(z.read(part));assert etree.tostring(changed.find('.//p:spTree',wm.NS))==etree.tostring(root.find('.//p:spTree',wm.NS))
