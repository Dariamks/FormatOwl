import sys
from pathlib import Path
import zipfile

import pytest
from PIL import ImageChops
from lxml import etree
from pypdf import PdfReader
from pptx import Presentation
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'scripts'))
from watermark_document_fixtures import build, MC
import watermark_file as wm


@pytest.fixture(scope='module')
def documents(tmp_path_factory):
    folder = tmp_path_factory.mktemp('documents')
    build(folder)
    return folder


def analyze(source, folder):
    return wm.analyze({'source': source, 'format': source.suffix[1:], 'folder': folder})


def output(source, data, folder, candidates=(), patches=None):
    path = folder/('out'+source.suffix)
    wm.export_file({'source': source, 'folder': folder, 'data': data,
                   'selection': {'candidates': list(candidates), 'regions': []},
                   'patches': patches or {}, 'output': path, 'prefix': 'out'})
    return path


def assert_unchanged_parts(source, result, except_parts):
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(result) as b:
        for name in a.namelist():
            if name not in except_parts: assert a.read(name) == b.read(name), name


def test_word_removes_every_compatibility_branch_and_preserves_body(documents, tmp_path):
    source = documents/'compatibility.docx'; data = analyze(source, tmp_path)
    marks = [c for c in data['candidates'] if 'WORD SAMPLE' in c['label']]
    assert len(marks) == 1
    out = output(source, data, tmp_path, [marks[0]['id']])
    with zipfile.ZipFile(out) as z:
        assert b'WORD SAMPLE' not in z.read(marks[0]['locator']['part'])
    assert_unchanged_parts(source, out, {marks[0]['locator']['part']})
    doc = Document(out)
    assert doc.tables[0].cell(1, 1).text == '42'
    assert doc.paragraphs[1].runs[0].bold
    assert 'KEEP normal header' in doc.sections[0].header.paragraphs[0].text
    assert 'CONFIDENTIAL' in doc.sections[0].footer.paragraphs[0].text
    assert len(data['pages']) == 2


def test_word_plain_footer_watermark_keeps_normal_header(documents, tmp_path):
    source = documents/'compatibility.docx'; data = analyze(source, tmp_path)
    mark = next(c for c in data['candidates'] if c['label'] == 'CONFIDENTIAL')
    assert mark['reason'] == 'header' and mark['pages'] == [0, 1]
    out = output(source, data, tmp_path, [mark['id']])
    doc = Document(out)
    assert all('CONFIDENTIAL' not in p.text for p in doc.sections[0].footer.paragraphs)
    assert 'KEEP normal header' in doc.sections[0].header.paragraphs[0].text
    assert_unchanged_parts(source, out, {mark['locator']['part']})


def test_word_image_repair_updates_both_alternatives_only_at_selected_reference(documents, tmp_path):
    source = documents/'image-compatibility.docx'; data = analyze(source, tmp_path)
    assert len(data['targets']) == 2, 'Choice and Fallback are one visual image, not two targets'
    target = data['targets'][0]
    patch = wm.image_open(tmp_path/target['file']); patch.paste('#183455', (255, 188, 358, 240)); patch.save(tmp_path/'patch.png')
    out = output(source, data, tmp_path, patches={target['id']: str(tmp_path/'patch.png')})
    with zipfile.ZipFile(out) as z:
        root = wm.read_xml(z.read('word/document.xml'))
        alternatives = root.xpath('//mc:AlternateContent//a:blip/@r:embed', namespaces={**wm.NS, 'mc': MC})
        assert len(alternatives) == 2 and len(set(alternatives)) == 1
        ordinary = root.xpath('//a:blip[not(ancestor::mc:AlternateContent)]/@r:embed', namespaces={**wm.NS, 'mc': MC})
        assert ordinary[0] != alternatives[0]
    assert_unchanged_parts(source, out, {'word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'})


def test_powerpoint_group_removes_plate_and_text_preserving_charts_tables_notes(documents, tmp_path):
    source = documents/'grouped.pptx'; data = analyze(source, tmp_path)
    mark = next(c for c in data['candidates'] if c['label'] == 'GROUP SAMPLE')
    root = wm.read_xml(wm.office_load(source)[mark['locator']['part']])
    node = root.xpath(mark['locator']['path'])[0]
    assert etree.QName(node).localname == 'grpSp', 'The plate and text must stay one candidate'
    assert mark['pages'] == [0]
    out = output(source, data, tmp_path, [mark['id']])
    prs = Presentation(out)
    assert len(prs.slides) == 2
    assert not any(s.shape_type == 6 for s in prs.slides[0].shapes)
    assert next(s for s in prs.slides[0].shapes if s.has_table).table.cell(1, 1).text == '42'
    assert next(s for s in prs.slides[0].shapes if s.has_chart).chart.series[0].values == (3., 7.)
    assert 'KEEP speaker notes' in prs.slides[0].notes_slide.notes_text_frame.text
    assert_unchanged_parts(source, out, {'ppt/slides/slide1.xml'})
    before = wm.image_open(tmp_path/'page-1.png'); after = wm.image_open(tmp_path/'out-1.png')
    assert not ImageChops.difference(before, after).convert('RGB').getbbox()


def test_powerpoint_hidden_slides_keep_page_mapping(documents, tmp_path):
    source = documents/'hidden.pptx'; data = analyze(source, tmp_path)
    assert len(data['pages']) == 2
    mark = next(c for c in data['candidates'] if c['label'] == 'GROUP SAMPLE')
    assert mark['pages'] == [0]
    out = output(source, data, tmp_path, [mark['id']])
    assert Presentation(out).slides[0]._element.get('show') == '0'


def test_powerpoint_repair_preserves_crop_rotation_and_other_shared_image(documents, tmp_path):
    source = documents/'grouped.pptx'; data = analyze(source, tmp_path)
    targets = [t for t in data['targets'] if t['locator']['part'] == 'ppt/slides/slide1.xml']
    assert len(targets) == 2
    target = targets[0]; patch = wm.image_open(tmp_path/target['file']); patch.paste('red', (260, 190, 355, 236)); patch.save(tmp_path/'patch.png')
    out = output(source, data, tmp_path, patches={target['id']: str(tmp_path/'patch.png')})
    before = Presentation(source); after = Presentation(out)
    old = [s for s in before.slides[0].shapes if s.shape_type == 13]
    new = [s for s in after.slides[0].shapes if s.shape_type == 13]
    assert new[0].image.blob != old[0].image.blob and new[1].image.blob == old[1].image.blob
    for a, b in zip(old, new):
        assert (a.crop_left, a.crop_right, a.rotation, a.width, a.height) == (b.crop_left, b.crop_right, b.rotation, b.width, b.height)
    assert_unchanged_parts(source, out, {'ppt/slides/slide1.xml', 'ppt/slides/_rels/slide1.xml.rels', '[Content_Types].xml'})


def test_pdf_plain_text_removal_keeps_shared_font_state_forms_links_and_bookmarks(documents, tmp_path):
    source = documents/'text-state.pdf'; data = analyze(source, tmp_path)
    mark = next(c for c in data['candidates'] if c['label'] == 'PDF SAMPLE')
    out = output(source, data, tmp_path, [mark['id']])
    before = PdfReader(source); after = PdfReader(out)
    text = after.pages[0].extract_text()
    assert 'PDF SAMPLE' not in text and 'KEEP body text' in text and 'KEEP shared font' in text
    assert after.get_fields()['keep']['/V'] == 'Kept'
    assert len(after.outline) == len(before.outline)
    assert str(after.pages[0]['/Annots'][0].get_object()['/A']['/URI']) == 'https://example.com'
    a = wm.image_open(tmp_path/'page-0.png'); b = wm.image_open(tmp_path/'out-0.png')
    assert a.crop((0, 760, a.width, a.height)).tobytes() == b.crop((0, 760, b.width, b.height)).tobytes()


@pytest.mark.parametrize('show', [b'(SAMPLE) Tj', b'[(SAM) 15 (PLE)] TJ', b'(SAMPLE) \'', b'1 2 (SAMPLE) "'])
def test_pdf_text_operators_retain_font_and_spacing_state(documents, tmp_path, show):
    import pikepdf
    with pikepdf.open(documents/'text-state.pdf') as pdf:
        page=pdf.pages[0]
        page.Contents=pdf.make_stream(b'BT /F1 28 Tf 3 Tc 4 Tw 10 50 Td '+show+b' ET BT 40 200 Td (KEEP) Tj ET')
        marks=wm.pdf_candidates(page,0)
        assert len(marks)==1
        wm.pdf_edit_container(pdf,page,[marks[0]['locator']],[])
        commands=wm.pdf_instructions(page)
        assert b'SAMPLE' not in page.Contents.read_bytes()
        assert b'(KEEP)' in page.Contents.read_bytes()
        assert [str(c.operator) for c in commands][:4]==['BT','Tf','Tc','Tw']
        path=tmp_path/'state.pdf';pdf.save(path)
    assert 'KEEP' in PdfReader(path).pages[0].extract_text()


def test_pdf_clipping_text_is_not_offered_as_independent_removal(documents):
    import pikepdf
    with pikepdf.open(documents/'text-state.pdf') as pdf:
        page=pdf.pages[0]
        page.Contents=pdf.make_stream(b'BT /F1 28 Tf 7 Tr 10 50 Td (SAMPLE) Tj ET 0 0 300 300 re f')
        assert wm.pdf_candidates(page,0)==[]


def test_powerpoint_normal_content_group_is_not_combined_with_watermark():
    root=wm.read_xml(f'''<p:spTree xmlns:p="{wm.NS['p']}" xmlns:a="{wm.NS['a']}">
      <p:grpSp><p:sp><p:txBody><a:p><a:r><a:t>KEEP body</a:t></a:r></a:p></p:txBody></p:sp>
      <p:sp><p:txBody><a:p><a:r><a:t>SAMPLE</a:t></a:r></a:p></p:txBody></p:sp></p:grpSp></p:spTree>'''.encode())
    nodes=wm.office_nodes(root,'pptx','ppt/slides/slide1.xml')
    assert len(nodes)==2 and all(etree.QName(n).localname=='sp' for n in nodes)
